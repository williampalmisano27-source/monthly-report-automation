import csv
from copy import copy
from datetime import date, datetime
from hashlib import sha256

import pytest
from openpyxl import load_workbook

from report_rollover.models import ValidationError
from report_rollover.rollover import next_month, roll_batch


@pytest.fixture
def batch(tmp_path, make_report):
    source = tmp_path / "input"
    source.mkdir()
    make_report(source / "report.xlsx")
    return source, tmp_path / "output"


def test_rolls_values_preserves_formatting_and_input(batch):
    source, output = batch
    path = source / "report.xlsx"
    before = sha256(path.read_bytes()).hexdigest()
    assert roll_batch(source, output) == output
    wb = load_workbook(output / "report.xlsx")
    original = load_workbook(path)
    try:
        ws, old = wb["Materials"], original["Materials"]
        assert ws["B2"].value == datetime(2026, 10, 1)
        assert [ws[c].value for c in ("D6", "E6", "F6", "G6")] == [85, 0, 0, "=D6+E6-F6"]
        assert [ws[c].value for c in ("A6", "B6", "C6")] == ["MAT-001", "Example material", "kg"]
        for attr in ("font", "fill", "border", "alignment", "number_format"):
            assert copy(getattr(ws["D6"], attr)) == copy(getattr(old["D6"], attr))
        assert ws.column_dimensions["B"].width == 28
        assert str(ws.merged_cells) == "A1:G1"
        assert ws.page_setup == old.page_setup
        assert ws.print_area == old.print_area
        assert ws.print_title_rows == old.print_title_rows
        assert wb.calculation.fullCalcOnLoad
    finally:
        wb.close()
        original.close()
    assert sha256(path.read_bytes()).hexdigest() == before
    with (output / "rollover-summary.csv").open(newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert rows == [{"source_filename": "report.xlsx", "location": "LOC-01", "material_id": "MAT-001",
                     "previous_month": "2026-09-01", "new_month": "2026-10-01",
                     "previous_closing": "85", "new_opening": "85"}]


@pytest.mark.parametrize("before,after", [
    (date(2026, 12, 1), date(2027, 1, 1)),
    (date(2024, 1, 1), date(2024, 2, 1)),
    (date(2024, 2, 1), date(2024, 3, 1)),
])
def test_calendar_boundary(before, after):
    assert next_month(before) == after


def test_final_calendar_month_is_actionable_error():
    with pytest.raises(ValidationError):
        next_month(date(9999, 12, 1))


@pytest.mark.parametrize("mode", ["existing", "nested", "same", "symlink", "dangling_symlink", "missing_parent"])
def test_refuses_unsafe_destinations(batch, mode):
    source, output = batch
    if mode == "existing":
        output.mkdir()
        (output / "keep.txt").write_text("keep")
    elif mode == "nested":
        output = source / "new"
    elif mode == "same":
        output = source
    elif mode == "symlink":
        output.symlink_to(source, target_is_directory=True)
    elif mode == "dangling_symlink":
        output.symlink_to(source / "absent", target_is_directory=True)
    elif mode == "missing_parent":
        output = output / "child"
    original = (source / "report.xlsx").read_bytes()
    with pytest.raises(ValidationError):
        roll_batch(source, output)
    assert (source / "report.xlsx").read_bytes() == original
    if mode == "existing":
        assert (output / "keep.txt").read_text() == "keep"


def test_bad_file_prevents_partial_publication(batch):
    source, output = batch
    original = (source / "report.xlsx").read_bytes()
    (source / "bad.xlsx").write_bytes(b"broken")
    with pytest.raises(ValidationError):
        roll_batch(source, output)
    assert not output.exists()
    assert not list(output.parent.glob(".rollover-*"))
    assert (source / "report.xlsx").read_bytes() == original


def test_failed_save_publishes_nothing(batch, monkeypatch):
    source, output = batch
    def fail_save(*args, **kwargs):
        raise OSError("simulated disk failure")
    monkeypatch.setattr("openpyxl.workbook.workbook.Workbook.save", fail_save)
    with pytest.raises(OSError, match="simulated disk failure"):
        roll_batch(source, output)
    assert not output.exists()
    assert not list(output.parent.glob(".rollover-*"))


@pytest.mark.parametrize("identifier", ["=1+1", "+1+1", "-1+1", "@SUM(1)", "\tformula", "\rformula", "\nformula"])
def test_summary_text_cannot_be_opened_as_formula(batch, identifier):
    source, output = batch
    path = source / "report.xlsx"
    wb = load_workbook(path)
    wb["Materials"]["A6"] = identifier
    wb["Materials"]["A6"].data_type = "s"
    wb.save(path)
    wb.close()
    roll_batch(source, output)
    with (output / "rollover-summary.csv").open(newline="") as stream:
        row = next(csv.DictReader(stream))
    assert row["material_id"].startswith("'")


def test_decimal_balance_survives_save(batch):
    source, output = batch
    path = source / "report.xlsx"
    wb = load_workbook(path)
    ws = wb["Materials"]
    ws["D6"], ws["E6"], ws["F6"] = 0.1, 0.2, 0.1
    wb.save(path)
    wb.close()
    roll_batch(source, output)
    wb = load_workbook(output / "report.xlsx")
    assert wb["Materials"]["D6"].value == 0.2
    wb.close()


@pytest.mark.parametrize("opening,received", [(9999999999999, 0.999), (1e308, 1e308), (1234567890123456, 7)])
def test_rejects_closing_that_excel_cannot_store_exactly(batch, opening, received):
    source, output = batch
    path = source / "report.xlsx"
    wb = load_workbook(path)
    ws = wb["Materials"]
    ws["D6"], ws["E6"], ws["F6"] = opening, received, 0
    wb.save(path)
    wb.close()
    before = path.read_bytes()
    with pytest.raises(ValidationError, match="(?i)(precision|exact|represent)"):
        roll_batch(source, output)
    assert not output.exists()
    assert not list(output.parent.glob(".rollover-*"))
    assert path.read_bytes() == before
