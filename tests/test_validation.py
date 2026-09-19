from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from openpyxl.styles import Font

from report_rollover.models import ValidationError
from report_rollover.validation import quantity, read_report, validate_batch


def change(source, action):
    wb = load_workbook(source)
    try:
        action(wb)
        wb.save(source)
    finally:
        wb.close()


def test_reads_balance_without_formula_cache(tmp_path, make_report):
    report = read_report(make_report(tmp_path / "report.xlsx"))
    assert report.rows[0].closing == Decimal("85")
    assert report.month == date(2026, 9, 1)
    assert report.location == "LOC-01"


@pytest.mark.parametrize("cell,value", [
    ("A6", None), ("B6", " "), ("C6", None), ("D6", True), ("D6", -1),
    ("E6", "25"), ("F6", 200), ("D6", 1.2345), ("G6", "=SUM(D6:F6)"),
    ("G6", 85), ("D6", "=1+1"), ("B2", "2026-09-01"),
    ("B2", date(2026, 9, 2)), ("B3", ""), ("A5", "Wrong header"),
])
def test_rejects_invalid_cells(tmp_path, make_report, cell, value):
    source = make_report(tmp_path / "bad.xlsx")
    change(source, lambda wb: setattr(wb["Materials"][cell], "value", value))
    with pytest.raises(ValidationError, match="bad.xlsx"):
        read_report(source)


@pytest.mark.parametrize("value", [float("inf"), float("nan"), Decimal("-Infinity"), True])
def test_rejects_nonfinite_or_boolean_quantities(value):
    with pytest.raises(ValidationError):
        quantity(value, "bad.xlsx:D6")


def test_rejects_missing_sheet(tmp_path, make_report):
    source = make_report(tmp_path / "bad.xlsx")
    change(source, lambda wb: setattr(wb.active, "title", "Wrong"))
    with pytest.raises(ValidationError, match="Materials"):
        read_report(source)


@pytest.mark.parametrize("kind", ["duplicate", "empty", "interior_blank"])
def test_rejects_broken_row_structure(tmp_path, make_report, kind):
    source = make_report(tmp_path / "bad.xlsx")
    def mutate(wb):
        ws = wb["Materials"]
        if kind == "empty":
            ws.delete_rows(6)
        else:
            row = 7 if kind == "duplicate" else 8
            values = ["MAT-001" if kind == "duplicate" else "MAT-002", "Other", "kg", 10, 1, 2, f"=D{row}+E{row}-F{row}"]
            for col, value in enumerate(values, 1):
                ws.cell(row, col, value)
    change(source, mutate)
    with pytest.raises(ValidationError):
        read_report(source)


def test_ignores_styled_blank_trailing_rows(tmp_path, make_report):
    source = make_report(tmp_path / "good.xlsx")
    change(source, lambda wb: setattr(wb["Materials"]["D100"], "font", Font(bold=True)))
    assert len(read_report(source).rows) == 1


def test_decimal_calculation_is_exact(tmp_path, make_report):
    source = make_report(tmp_path / "good.xlsx")
    def mutate(wb):
        for cell, value in [("D6", 0.1), ("E6", 0.2), ("F6", 0.1)]:
            wb["Materials"][cell] = value
    change(source, mutate)
    assert read_report(source).rows[0].closing == Decimal("0.2")


def test_corrupt_workbook_has_filename(tmp_path):
    source = tmp_path / "broken.xlsx"
    source.write_bytes(b"not a workbook")
    with pytest.raises(ValidationError, match="broken.xlsx"):
        read_report(source)


def test_batch_is_sorted_and_materials_can_repeat_across_locations(tmp_path, make_report):
    make_report(tmp_path / "b.xlsx", location="B")
    make_report(tmp_path / "a.xlsx", location="A")
    (tmp_path / "~$a.xlsx").write_bytes(b"lock")
    assert [r.location for r in validate_batch(tmp_path)] == ["A", "B"]


@pytest.mark.parametrize("kind", ["empty", "duplicate_location", "mixed_month", "bad_file", "missing_directory"])
def test_rejects_invalid_batches(tmp_path, make_report, kind):
    if kind == "missing_directory":
        tmp_path = tmp_path / "absent"
    elif kind != "empty":
        make_report(tmp_path / "a.xlsx")
        if kind == "bad_file":
            (tmp_path / "b.xlsx").write_bytes(b"broken")
        else:
            make_report(tmp_path / "b.xlsx", location="OTHER" if kind == "mixed_month" else "LOC-01",
                        month=date(2026, 10, 1) if kind == "mixed_month" else date(2026, 9, 1))
    with pytest.raises(ValidationError):
        validate_batch(tmp_path)


@pytest.mark.parametrize("kind", ["style_index", "shared_string", "font_size"])
def test_structural_corruption_is_file_specific(tmp_path, make_report, corrupt_workbook, kind):
    source = make_report(tmp_path / "malformed.xlsx")
    corrupt_workbook(source, kind)
    with pytest.raises(ValidationError, match="malformed.xlsx"):
        read_report(source)


def test_literal_formula_text_is_not_a_calculating_cell(tmp_path, make_report):
    source = make_report(tmp_path / "literal.xlsx")
    change(source, lambda wb: setattr(wb["Materials"]["G6"], "data_type", "s"))
    with pytest.raises(ValidationError, match="literal.xlsx:G6"):
        read_report(source)
