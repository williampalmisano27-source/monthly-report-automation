import csv
from datetime import datetime
import subprocess
import sys

import pytest
from openpyxl import load_workbook

from report_rollover.models import ValidationError
from report_rollover.rollover import roll_batch
from report_rollover.samples import generate_samples
from report_rollover.validation import read_report


def test_twelve_reports_produce_96_correct_audit_rows(tmp_path):
    demo = tmp_path / "demo"
    generate_samples(demo)
    inputs = sorted((demo / "input").glob("*.xlsx"))
    assert len(inputs) == 12
    assert len(list((demo / "invalid").glob("*.xlsx"))) == 3
    assert len({read_report(p).location for p in inputs}) == 12
    roll_batch(demo / "input", demo / "output")
    outputs = sorted((demo / "output").glob("*.xlsx"))
    assert len(outputs) == 12
    with (demo / "output" / "rollover-summary.csv").open(newline="") as stream:
        audit = list(csv.DictReader(stream))
    assert len(audit) == 96
    for i, path in enumerate(outputs, 1):
        wb = load_workbook(path)
        ws = wb["Materials"]
        assert ws["B2"].value == datetime(2026, 10, 1)
        assert ws.max_row == 13
        assert str(ws.merged_cells) == "A1:G1"
        assert ws.page_setup.orientation == "landscape"
        for j in range(1, 9):
            row = j + 5
            expected = 115 + 9 * i + 2 * j
            assert ws.cell(row, 4).value == expected
            assert ws.cell(row, 5).value == 0
            assert ws.cell(row, 6).value == 0
            assert ws.cell(row, 7).value == f"=D{row}+E{row}-F{row}"
            audit_row = audit[(i - 1) * 8 + j - 1]
            assert audit_row["new_opening"] == str(expected)
            assert audit_row["previous_closing"] == str(expected)
        wb.close()


def test_invalid_samples_explain_their_error(tmp_path):
    demo = tmp_path / "demo"
    generate_samples(demo)
    for filename, fragment in [("duplicate-id.xlsx", "duplicate"), ("missing-id.xlsx", "A6"), ("invalid-quantity.xlsx", "numeric")]:
        with pytest.raises(ValidationError, match=fragment):
            read_report(demo / "invalid" / filename)


def test_generator_refuses_existing_directory(tmp_path):
    (tmp_path / "keep.txt").write_text("original")
    with pytest.raises(ValidationError):
        generate_samples(tmp_path)
    assert (tmp_path / "keep.txt").read_text() == "original"


def test_generator_cli(tmp_path):
    result = subprocess.run([sys.executable, "-m", "report_rollover.samples", "--output", str(tmp_path / "demo")], capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert len(list((tmp_path / "demo" / "input").glob("*.xlsx"))) == 12
