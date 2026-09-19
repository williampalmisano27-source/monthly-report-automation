from datetime import date

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


@pytest.fixture
def make_report():
    def create(path, *, month=date(2026, 9, 1), location="LOC-01"):
        wb = Workbook()
        ws = wb.active
        ws.title = "Materials"
        ws.merge_cells("A1:G1")
        ws["A1"] = "Synthetic material report"
        ws["B2"] = month
        ws["B3"] = location
        ws.append([])
        for column, text in enumerate(
            ["Material ID", "Description", "Unit", "Opening", "Received", "Used", "Closing"], 1
        ):
            ws.cell(5, column, text)
        for column, value in enumerate(
            ["MAT-001", "Example material", "kg", 100, 25, 40, "=D6+E6-F6"], 1
        ):
            ws.cell(6, column, value)
        ws["D6"].font = Font(name="Calibri", size=12, color="2456A6")
        ws["D6"].fill = PatternFill("solid", fgColor="EAF2F8")
        ws["D6"].border = Border(bottom=Side(style="thin", color="CCDDEE"))
        ws["D6"].alignment = Alignment(horizontal="right")
        ws["D6"].number_format = "0.000"
        ws.column_dimensions["B"].width = 28
        ws.page_setup.orientation = "landscape"
        ws.print_area = "A1:G6"
        ws.print_title_rows = "1:5"
        wb.save(path)
        wb.close()
        return path
    return create



@pytest.fixture
def corrupt_workbook():
    def corrupt(path, kind):
        import re
        from zipfile import ZipFile
        with ZipFile(path) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        sheet_name = "xl/worksheets/sheet1.xml"
        if kind == "style_index":
            text = entries[sheet_name].decode()
            text = re.sub(r'(<c r="D6"[^>]*?) s="\d+"', r'\1 s="9999"', text)
            entries[sheet_name] = text.encode()
        elif kind == "shared_string":
            text = entries[sheet_name].decode()
            text = re.sub(r'<c r="A6"[^>]*>.*?</c>', '<c r="A6" t="s"><v>9999</v></c>', text)
            entries[sheet_name] = text.encode()
        elif kind == "font_size":
            entries["xl/styles.xml"] = entries["xl/styles.xml"].replace(b'<sz val="12"', b'<sz val="garbage"')
        else:
            raise AssertionError(kind)
        with ZipFile(path, "w") as archive:
            for name, content in entries.items():
                archive.writestr(name, content)
    return corrupt
