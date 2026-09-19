"""Publish complete, validated output batches while leaving source files intact."""
import csv
from datetime import date
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook
from openpyxl.workbook.properties import CalcProperties

from .models import ValidationError
from .validation import excel_number, read_report, validate_batch


def next_month(month: date) -> date:
    try:
        return date(month.year + (month.month == 12), month.month % 12 + 1, 1)
    except ValueError as exc:
        raise ValidationError("Cannot advance beyond December 9999") from exc


def _csv_text(value: str) -> str:
    # A CSV quote only escapes CSV syntax, not spreadsheet formula evaluation.
    if value.startswith(("=", "+", "-", "@", "\t", "\r", "\n")) or value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _destination(input_dir: Path, output_dir: Path) -> Path:
    if output_dir.is_symlink() or output_dir.exists():
        raise ValidationError(f"{output_dir}: destination already exists; choose a new directory")
    source = input_dir.resolve()
    destination = output_dir.resolve()
    if destination == source or source in destination.parents:
        raise ValidationError("Output directory must be outside the input directory")
    if not destination.parent.is_dir():
        raise ValidationError(f"{destination.parent}: output parent directory must already exist")
    return destination


def roll_batch(input_dir: Path, output_dir: Path) -> Path:
    input_dir, output_dir = Path(input_dir), Path(output_dir)
    destination = _destination(input_dir, output_dir)
    reports = validate_batch(input_dir)
    month = next_month(reports[0].month)
    staging = Path(tempfile.mkdtemp(prefix=".rollover-", dir=destination.parent))
    try:
        with (staging / "rollover-summary.csv").open("w", newline="", encoding="utf-8") as stream:
            writer = csv.writer(stream)
            writer.writerow(["source_filename", "location", "material_id", "previous_month",
                             "new_month", "previous_closing", "new_opening"])
            for report in reports:
                wb = load_workbook(report.source, data_only=False, keep_links=False)
                try:
                    ws = wb["Materials"]
                    ws["B2"] = month
                    for row in report.rows:
                        ws.cell(row.row_number, 4).value = excel_number(
                            row.closing, f"{report.source.name}:G{row.row_number}")
                        ws.cell(row.row_number, 5).value = 0
                        ws.cell(row.row_number, 6).value = 0
                        writer.writerow([_csv_text(report.source.name), _csv_text(report.location),
                                         _csv_text(row.material_id), report.month.isoformat(), month.isoformat(),
                                         format(row.closing, "f"), format(row.closing, "f")])
                    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True, forceFullCalc=True)
                    wb.save(staging / report.source.name)
                finally:
                    wb.close()
                saved = read_report(staging / report.source.name)
                for expected, actual in zip(report.rows, saved.rows, strict=True):
                    if actual.opening != expected.closing:
                        raise ValidationError(
                            f"{report.source.name}:D{actual.row_number}: "
                            "calculated closing could not be stored exactly in Excel")
        # One local writer is assumed. This recheck prevents ordinary rerun overwrites;
        # it is not a lock against adversarial concurrent filesystem changes.
        _destination(input_dir, output_dir)
        staging.rename(destination)
    finally:
        if staging.exists():
            shutil.rmtree(staging)
    return destination
