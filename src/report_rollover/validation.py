"""Read the documented template without trusting cached formula results."""
from datetime import date, datetime
from decimal import Decimal, DecimalException, localcontext
import math
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import MaterialRow, Report, ValidationError

HEADERS = ("Material ID", "Description", "Unit", "Opening", "Received", "Used", "Closing")


def excel_number(value: Decimal, label: str) -> float:
    """Reject quantities that cannot retain their precision as Excel numbers."""
    converted = float(value)
    if not math.isfinite(converted) or Decimal(format(converted, ".15g")) != value:
        raise ValidationError(f"{label}: quantity exceeds exact Excel precision (15 significant digits)")
    return converted


def quantity(value, label: str) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise ValidationError(f"{label}: expected a numeric quantity, not text or a formula")
    try:
        number = Decimal(str(value))
        if not number.is_finite() or number < 0:
            raise ValidationError(f"{label}: quantity must be finite and nonnegative")
        with localcontext() as context:
            context.prec = 400
            if number != number.quantize(Decimal("0.001")):
                raise ValidationError(f"{label}: maximum three decimal places")
    except DecimalException as exc:
        raise ValidationError(f"{label}: unsupported numeric precision") from exc
    excel_number(number, label)
    return number


def _text(cell, path: Path) -> str:
    value = cell.value
    if cell.data_type == "f" or not isinstance(value, str) or not value.strip():
        raise ValidationError(f"{path.name}:{cell.coordinate}: expected nonempty literal text")
    return value


def read_report(path: Path) -> Report:
    path = Path(path)
    try:
        wb = load_workbook(path, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException, ParseError, KeyError, ValueError, OSError,
            EOFError, IndexError, TypeError) as exc:
        raise ValidationError(f"{path.name}: cannot read workbook ({exc})") from exc
    try:
        if "Materials" not in wb.sheetnames:
            raise ValidationError(f"{path.name}: missing Materials sheet")
        ws = wb["Materials"]
        if tuple(ws.cell(5, col).value for col in range(1, 8)) != HEADERS:
            raise ValidationError(f"{path.name}: row 5 does not match required column headers")
        month = ws["B2"].value
        if isinstance(month, datetime):
            if month.time().isoformat() != "00:00:00":
                raise ValidationError(f"{path.name}:B2: use a date without a time")
            month = month.date()
        if not isinstance(month, date) or month.day != 1:
            raise ValidationError(f"{path.name}:B2: expected an Excel date on the first of the month")
        location = _text(ws["B3"], path)
        last_row = ws.max_row
        while last_row >= 6 and all(ws.cell(last_row, col).value is None for col in range(1, 8)):
            last_row -= 1
        if last_row < 6:
            raise ValidationError(f"{path.name}: no material records")
        rows = []
        seen = set()
        for row in range(6, last_row + 1):
            material_id, description, unit = (_text(ws.cell(row, col), path) for col in range(1, 4))
            if material_id.strip() in seen:
                raise ValidationError(f"{path.name}:A{row}: duplicate material ID {material_id!r}")
            seen.add(material_id.strip())
            opening, received, used = (
                quantity(ws.cell(row, col).value, f"{path.name}:{ws.cell(row, col).coordinate}")
                for col in range(4, 7)
            )
            formula = f"=D{row}+E{row}-F{row}"
            if ws.cell(row, 7).data_type != "f" or ws.cell(row, 7).value != formula:
                raise ValidationError(f"{path.name}:G{row}: expected supported formula {formula}")
            with localcontext() as context:
                context.prec = 400
                closing = opening + received - used
            if closing < 0:
                raise ValidationError(f"{path.name}:G{row}: calculated closing balance is negative")
            excel_number(closing, f"{path.name}:G{row}")
            rows.append(MaterialRow(row, material_id, description, unit, opening, received, used, closing))
        return Report(path, month, location, tuple(rows))
    finally:
        wb.close()


def validate_batch(input_dir: Path) -> tuple[Report, ...]:
    input_dir = Path(input_dir)
    if not input_dir.is_dir():
        raise ValidationError(f"{input_dir}: input must be an existing directory")
    paths = sorted(p for p in input_dir.iterdir()
                   if p.is_file() and p.suffix.lower() == ".xlsx" and not p.name.startswith("~$"))
    if not paths:
        raise ValidationError(f"{input_dir}: no .xlsx workbooks found")
    reports = []
    errors = []
    for path in paths:
        try:
            reports.append(read_report(path))
        except ValidationError as exc:
            errors.append(str(exc))
    locations = set()
    for report in reports:
        if report.location.strip() in locations:
            errors.append(f"{report.source.name}:B3: duplicate location {report.location!r}")
        locations.add(report.location.strip())
        if report.month != reports[0].month:
            errors.append(f"{report.source.name}:B2: mixed reporting months in batch")
    if errors:
        raise ValidationError("Batch validation failed:\n" + "\n".join(errors))
    return tuple(reports)
