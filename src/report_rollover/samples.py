"""Generate the disclosed synthetic portfolio fixtures. No client data is used."""
import argparse
from datetime import date
from pathlib import Path
import shutil
import sys
import tempfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .models import ValidationError
from .validation import HEADERS

MATERIALS = ("Aggregate", "Cement", "Sand", "Lime", "Gypsum", "Clay", "Gravel", "Pigment")


def _sample(path: Path, location: int) -> None:
    wb = Workbook()
    try:
        ws = wb.active
        ws.title = "Materials"
        ws.sheet_view.showGridLines = False
        ws.merge_cells("A1:G1")
        ws["A1"] = "MATERIAL REPORT | SYNTHETIC PORTFOLIO DEMO"
        ws["A1"].font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        ws["A1"].fill = PatternFill("solid", fgColor="173F46")
        ws["A1"].alignment = Alignment(vertical="center")
        ws.row_dimensions[1].height = 38
        ws["A2"], ws["B2"] = "Report month", date(2026, 9, 1)
        ws["B2"].number_format = "mmmm yyyy"
        ws["A3"], ws["B3"] = "Location", f"LOC-{location:02d}"
        ws["D3"] = "All quantities in kg"
        ws["D3"].font = Font(name="Calibri", size=10, italic=True, color="536A70")
        for col, heading in enumerate(HEADERS, 1):
            cell = ws.cell(5, col, heading)
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="286570")
            cell.alignment = Alignment(vertical="center", horizontal="left" if col < 4 else "right")
        ws.row_dimensions[5].height = 28
        for j, description in enumerate(MATERIALS, 1):
            row = j + 5
            values = [f"MAT-{j:03d}", description, "kg", 100 + 10 * location + j,
                      20 + j, 5 + location, f"=D{row}+E{row}-F{row}"]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row, col, value)
                cell.font = Font(name="Calibri", size=11, color="2456A6" if 4 <= col <= 6 else "193A40",
                                 bold=col == 7)
                cell.fill = PatternFill("solid", fgColor="EDF5F4" if row % 2 == 0 else "FFFFFF")
                cell.border = Border(bottom=Side(style="hair", color="D5E3E1"))
                cell.alignment = Alignment(vertical="center", horizontal="right" if col >= 4 else "left")
                if col >= 4:
                    cell.number_format = '#,##0.000;[Red](#,##0.000);"–"'
            ws.row_dimensions[row].height = 25
        for col, width in {"A": 19, "B": 27, "C": 10, "D": 17, "E": 17, "F": 17, "G": 17}.items():
            ws.column_dimensions[col].width = width
        ws.freeze_panes = "D6"
        ws.auto_filter.ref = "A5:G13"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_area = "A1:G13"
        ws.print_title_rows = "1:5"
        ws.oddFooter.center.text = "Independent portfolio project | Synthetic data"
        wb.save(path)
    finally:
        wb.close()


def generate_samples(output: Path) -> None:
    output = Path(output)
    if output.exists() or output.is_symlink():
        raise ValidationError(f"{output}: sample destination already exists")
    if not output.parent.is_dir():
        raise ValidationError(f"{output.parent}: parent directory must already exist")
    staging = Path(tempfile.mkdtemp(prefix=".samples-", dir=output.parent))
    try:
        (staging / "input").mkdir()
        (staging / "invalid").mkdir()
        for location in range(1, 13):
            _sample(staging / "input" / f"location-{location:02d}.xlsx", location)
        for filename, cell, value in [("duplicate-id.xlsx", "A7", "MAT-001"),
                                      ("missing-id.xlsx", "A6", None),
                                      ("invalid-quantity.xlsx", "D6", "unknown")]:
            wb = load_workbook(staging / "input" / "location-01.xlsx")
            try:
                wb["Materials"][cell] = value
                wb.save(staging / "invalid" / filename)
            finally:
                wb.close()
        if output.exists() or output.is_symlink():
            raise ValidationError(f"{output}: sample destination already exists")
        staging.rename(output)
    finally:
        if staging.exists():
            shutil.rmtree(staging)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Create 12 synthetic reports and 3 invalid examples.")
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        generate_samples(args.output)
    except ValidationError as exc:
        print(f"Validation error: {exc}", file=sys.stderr)
        return 2
    except OSError as exc:
        print(f"File operation failed: {exc}", file=sys.stderr)
        return 1
    print(f"Created synthetic samples in {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
